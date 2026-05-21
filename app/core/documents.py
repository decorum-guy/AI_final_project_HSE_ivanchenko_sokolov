from __future__ import annotations

import re
from datetime import datetime
from html import escape, unescape
from io import BytesIO

from bs4 import BeautifulSoup, NavigableString, Tag

from app.core.gigachat_client import repair_digest_text
from app.db.models import DigestHistory


def _plain_text_from_html(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for link in soup.find_all("a"):
        href = link.get("href")
        if href:
            link.string = f"{link.get_text(strip=True)} ({href})"
    text = soup.get_text("\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return unescape(text).strip()


def _period_title(period: str) -> str:
    return {
        "today": "за сегодня",
        "3days": "за последние 3 дня",
        "week": "за неделю",
    }.get(period, period)


def _mode_title(source_mode: str) -> str:
    return {
        "interests": "по интересам",
        "selected_sources": "по выбранным источникам",
    }.get(source_mode, source_mode)


def _document_title(digest: DigestHistory) -> str:
    created = digest.created_at.strftime("%d.%m.%Y %H:%M") if isinstance(digest.created_at, datetime) else ""
    return f"ИнфоПульс: дайджест {_period_title(digest.period)} от {created}".strip()


def build_digest_docx(digest: DigestHistory) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    document = Document()
    document.add_heading(_document_title(digest), level=1)
    document.add_paragraph(f"Период: {_period_title(digest.period)}")
    document.add_paragraph(f"Режим: {_mode_title(digest.source_mode)}")
    document.add_paragraph("")

    def add_hyperlink(paragraph, text: str, url: str) -> None:
        relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)

        run = OxmlElement("w:r")
        run_properties = OxmlElement("w:rPr")

        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        run_properties.append(color)

        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        run_properties.append(underline)

        run.append(run_properties)
        text_element = OxmlElement("w:t")
        text_element.text = text
        run.append(text_element)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def add_inline(paragraph, node, *, bold: bool = False, italic: bool = False) -> None:
        if isinstance(node, NavigableString):
            if str(node):
                text_run = paragraph.add_run(str(node))
                text_run.bold = bold
                text_run.italic = italic
            return
        if not isinstance(node, Tag):
            return

        if node.name == "a":
            href = node.get("href", "")
            text = node.get_text(strip=True) or href
            if href:
                add_hyperlink(paragraph, text, href)
            else:
                paragraph.add_run(text)
            return

        next_bold = bold or node.name == "b"
        next_italic = italic or node.name == "i"
        for child in node.children:
            add_inline(paragraph, child, bold=next_bold, italic=next_italic)

    def add_html_line(line: str, style=None) -> None:
        paragraph = document.add_paragraph(style=style)
        line_soup = BeautifulSoup(line, "html.parser")
        for node in line_soup.contents:
            add_inline(paragraph, node)

    html = repair_digest_text(digest.digest_text)
    html = html.replace("<blockquote>", "\n<blockquote>").replace("</blockquote>", "</blockquote>\n")
    for raw_line in html.splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph("")
            continue
        style = "Intense Quote" if line.lower().startswith("<blockquote>") else None
        add_html_line(line, style=style)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def build_digest_pdf(digest: DigestHistory) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    stream = BytesIO()
    font_name = "Helvetica"
    for font_path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ):
        try:
            pdfmetrics.registerFont(TTFont("DigestSans", font_path))
            font_name = "DigestSans"
            break
        except Exception:
            continue

    styles = getSampleStyleSheet()
    normal = ParagraphStyle("DigestNormal", parent=styles["Normal"], fontName=font_name, fontSize=10, leading=14)
    title = ParagraphStyle("DigestTitle", parent=styles["Title"], fontName=font_name, fontSize=16, leading=20)
    quote = ParagraphStyle(
        "DigestQuote",
        parent=normal,
        leftIndent=12,
        borderColor=colors.lightgrey,
        borderWidth=1,
        borderPadding=8,
        backColor=colors.whitesmoke,
    )

    doc = SimpleDocTemplate(
        stream,
        pagesize=A4,
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )
    story = [Paragraph(_document_title(digest), title), Spacer(1, 8)]
    story.append(Paragraph(f"Период: {_period_title(digest.period)}<br/>Режим: {_mode_title(digest.source_mode)}", normal))
    story.append(Spacer(1, 12))

    soup = BeautifulSoup(repair_digest_text(digest.digest_text), "html.parser")

    def add_pdf_text(text: str, style=normal) -> None:
        for line in text.splitlines():
            line = line.strip()
            if line:
                story.append(Paragraph(line, style))
            else:
                story.append(Spacer(1, 6))

    for child in soup.children:
        if isinstance(child, NavigableString):
            add_pdf_text(escape(str(child)))
            continue
        if not isinstance(child, Tag):
            continue
        if child.name == "blockquote":
            add_pdf_text(escape(child.get_text(" ", strip=True)), quote)
        else:
            add_pdf_text(str(child), normal)

    doc.build(story)
    return stream.getvalue()
